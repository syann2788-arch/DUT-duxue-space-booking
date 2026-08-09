import asyncio
import io
from datetime import date, timedelta

from openpyxl import load_workbook
from sqlalchemy import select
from conftest import auth_header
from app.database import async_session
from app.models import CleanupVerification, Reservation, ReservationStatus, Room, SceneType, UsageMode, User, UserRole


async def force_finished(reservation_id: int):
    async with async_session() as db:
        reservation = await db.get(Reservation, reservation_id)
        reservation.date = date.today() - timedelta(days=1)
        reservation.status = ReservationStatus.cleanup_pending
        await db.commit()


async def force_cleanup_review(reservation_ids: list[int]):
    async with async_session() as db:
        for reservation_id in reservation_ids:
            reservation = await db.get(Reservation, reservation_id)
            reservation.date = date.today() - timedelta(days=1)
            reservation.start_minute = 8 * 60
            reservation.end_minute = 8 * 60 + 30
            reservation.status = ReservationStatus.cleanup_pending
            db.add(CleanupVerification(
                reservation_id=reservation_id,
                photo_urls=[f"/uploads/cleanup_{reservation_id}.jpg"],
            ))
        await db.commit()


async def force_approved_past(reservation_id: int):
    async with async_session() as db:
        reservation = await db.get(Reservation, reservation_id)
        reservation.date = date.today() - timedelta(days=1)
        reservation.start_minute = 8 * 60
        reservation.end_minute = 8 * 60 + 30
        reservation.status = ReservationStatus.approved
        await db.commit()


async def force_used(reservation_id: int):
    async with async_session() as db:
        reservation = await db.get(Reservation, reservation_id)
        reservation.status = ReservationStatus.completed
        await db.commit()


async def set_user_role(student_id: str, role: UserRole):
    async with async_session() as db:
        user = (await db.execute(select(User).where(User.student_id == student_id))).scalar_one()
        user.role = role
        await db.commit()


async def create_fragmented_room_occupancy(day: date, student_id: str):
    """Leave a different study room free in each adjacent half-hour."""
    async with async_session() as db:
        user = (await db.execute(select(User).where(User.student_id == student_id))).scalar_one()
        rooms = {
            item.room_code: item
            for item in (await db.execute(select(Room).where(Room.room_code.in_(("A101", "A102", "A105"))))).scalars()
        }
        for room_code, start_slot, end_slot in (
            ("A101", 0, 1),
            ("A102", 1, 2),
            ("A105", 0, 2),
        ):
            db.add(Reservation(
                user_id=user.id,
                room_id=rooms[room_code].id,
                date=day,
                start_slot=start_slot,
                end_slot=end_slot,
                start_minute=8 * 60 + start_slot * 30,
                end_minute=8 * 60 + end_slot * 30,
                start_hour=8,
                end_hour=9,
                scene=SceneType.meeting,
                usage_mode=UsageMode.exclusive,
                people_count=1,
                purpose="测试相邻时段候选房间交叉占用",
                campus_card_photo_url="/uploads/campus_card_test.jpg",
                status=ReservationStatus.approved,
            ))
        await db.commit()


async def create_disjoint_shared_occupancy(day: date, student_id: str):
    """Two adjacent crowds must be evaluated by peak, not summed together."""
    async with async_session() as db:
        user = (await db.execute(select(User).where(User.student_id == student_id))).scalar_one()
        room = (await db.execute(select(Room).where(Room.room_code == "A102"))).scalar_one()
        for start_slot, end_slot in ((10, 11), (11, 12)):
            db.add(Reservation(
                user_id=user.id,
                room_id=room.id,
                date=day,
                start_slot=start_slot,
                end_slot=end_slot,
                start_minute=8 * 60 + start_slot * 30,
                end_minute=8 * 60 + end_slot * 30,
                start_hour=13,
                end_hour=14,
                scene=SceneType.study,
                usage_mode=UsageMode.shared,
                people_count=5,
                purpose="测试相邻共享预约峰值容量",
                campus_card_photo_url="/uploads/campus_card_test.jpg",
                status=ReservationStatus.approved,
            ))
        await db.commit()


async def create_legacy_unknown_mode_occupancy(day: date, student_id: str):
    async with async_session() as db:
        user = (await db.execute(select(User).where(User.student_id == student_id))).scalar_one()
        room = (await db.execute(select(Room).where(Room.room_code == "A102"))).scalar_one()
        db.add(Reservation(
            user_id=user.id,
            room_id=room.id,
            date=day,
            start_slot=14,
            end_slot=15,
            start_minute=15 * 60,
            end_minute=15 * 60 + 30,
            start_hour=15,
            end_hour=16,
            scene=None,
            usage_mode=None,
            people_count=1,
            purpose="第一版历史预约",
            campus_card_photo_url=None,
            status=ReservationStatus.approved,
        ))
        await db.commit()


def register(client, suffix: int) -> str:
    response = client.post("/api/auth/register", json={
        "student_id": f"202600{suffix:02d}",
        "name": f"测试学生{suffix}",
        "phone": f"139000000{suffix:02d}",
        "class_name": "2601",
        "password": "Test1234",
    })
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


def test_health_and_config(client):
    assert client.get("/api/health").json()["version"] == "2.0.0"
    config = client.get("/api/reservations/config").json()
    assert config["slot_minutes"] == 30
    assert config["max_minutes_per_day"] == 240
    numeric_only = client.post("/api/auth/register", json={
        "student_id": "20269999", "name": "密码规则测试", "phone": "13800009999",
        "class_name": "2601", "password": "123456",
    })
    assert numeric_only.status_code == 422
    assert "密码必须同时包含字母和数字" in numeric_only.text

    invalid_class = client.post("/api/auth/register", json={
        "student_id": "20269998", "name": "班级规则测试", "phone": "13800009998",
        "class_name": "笃学2601", "password": "Test1234",
    })
    assert invalid_class.status_code == 422
    assert "班级必须为4位数字" in invalid_class.text

    admin_login = client.post("/api/auth/login", json={"student_id": "admin001", "password": "admin123"})
    availability = client.get("/api/reservations/availability", headers=auth_header(admin_login.json()["access_token"]), params={
        "scene": "study",
        "date": (date.today() + timedelta(days=1)).isoformat(),
        "people_count": 1,
    })
    assert availability.status_code == 200, availability.text
    assert availability.json()["scene"] == "study"
    assert len(availability.json()["slots"]) == 28
    assert "available_room_ids" in availability.json()["slots"][0]


def test_permission_scoped_a106_and_continuous_room_availability(client):
    student_token = register(client, 41)
    counselor_student_id = "20260042"
    register(client, 42)
    asyncio.run(set_user_role(counselor_student_id, UserRole.counselor))
    counselor_login = client.post("/api/auth/login", json={
        "student_id": counselor_student_id,
        "password": "Test1234",
    })
    assert counselor_login.status_code == 200, counselor_login.text
    counselor_token = counselor_login.json()["access_token"]

    day = (date.today() + timedelta(days=7)).isoformat()
    meeting = {
        "scene": "meeting",
        "date": day,
        "start_slot": 22,
        "end_slot": 23,
        "people_count": 2,
        "purpose": "开展师生交流与工作沟通会议",
        "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    }
    student_booking = client.post("/api/reservations", headers=auth_header(student_token), json=meeting)
    assert student_booking.status_code == 201, student_booking.text
    assert student_booking.json()["room"]["room_code"] == "A105"

    counselor_booking = client.post("/api/reservations", headers=auth_header(counselor_token), json=meeting)
    assert counselor_booking.status_code == 201, counselor_booking.text
    assert counselor_booking.json()["room"]["room_code"] == "A106"

    rooms = client.get("/api/rooms").json()
    a106_id = next(item["id"] for item in rooms if item["room_code"] == "A106")
    query = {"scene": "meeting", "date": day, "people_count": 2}
    student_slots = client.get(
        "/api/reservations/availability", headers=auth_header(student_token), params=query,
    ).json()["slots"]
    counselor_slots = client.get(
        "/api/reservations/availability", headers=auth_header(counselor_token), params=query,
    ).json()["slots"]
    assert all(a106_id not in slot["available_room_ids"] for slot in student_slots)
    assert a106_id in counselor_slots[0]["available_room_ids"]

    fragmented_day = date.today() + timedelta(days=5)
    asyncio.run(create_fragmented_room_occupancy(fragmented_day, "20260041"))
    fragmented = client.get(
        "/api/reservations/availability",
        headers=auth_header(student_token),
        params={"scene": "study", "date": fragmented_day.isoformat(), "people_count": 1},
    )
    assert fragmented.status_code == 200, fragmented.text
    first, second = fragmented.json()["slots"][:2]
    assert first["available"] is True and second["available"] is True
    assert set(first["available_room_ids"]).isdisjoint(second["available_room_ids"])
    full_range = client.post("/api/reservations", headers=auth_header(student_token), json={
        "scene": "study",
        "date": fragmented_day.isoformat(),
        "start_slot": 0,
        "end_slot": 2,
        "people_count": 1,
        "purpose": "个人自习",
        "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    })
    assert full_range.status_code == 409

    capacity_token = register(client, 43)
    capacity_day = date.today() + timedelta(days=4)
    asyncio.run(create_disjoint_shared_occupancy(capacity_day, "20260041"))
    peak_capacity = client.post("/api/reservations", headers=auth_header(capacity_token), json={
        "scene": "study",
        "date": capacity_day.isoformat(),
        "start_slot": 10,
        "end_slot": 12,
        "people_count": 1,
        "purpose": "个人自习",
        "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    })
    assert peak_capacity.status_code == 201, peak_capacity.text
    assert peak_capacity.json()["room"]["room_code"] == "A102"

    legacy_token = register(client, 44)
    asyncio.run(create_legacy_unknown_mode_occupancy(capacity_day, "20260041"))
    legacy_conflict = client.post("/api/reservations", headers=auth_header(legacy_token), json={
        "scene": "study",
        "date": capacity_day.isoformat(),
        "start_slot": 14,
        "end_slot": 15,
        "people_count": 1,
        "purpose": "个人自习",
        "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    })
    assert legacy_conflict.status_code == 201, legacy_conflict.text
    assert legacy_conflict.json()["room"]["room_code"] == "A101"


def test_wechat_login_requires_server_credentials(client, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "WECHAT_APP_ID", "")
    monkeypatch.setattr(settings, "WECHAT_APP_SECRET", "")
    response = client.post("/api/auth/wechat/login", json={"code": "test-code"})
    assert response.status_code == 503
    assert "AppID/AppSecret" in response.json()["detail"]


def test_wechat_login_reports_unbound_account(client, monkeypatch):
    from app.routers import auth as auth_router

    async def fake_code2session(code: str) -> str:
        assert code == "unbound-code"
        return "unbound-openid"

    monkeypatch.setattr(auth_router, "_code2session", fake_code2session)
    response = client.post("/api/auth/wechat/login", json={"code": "unbound-code"})
    assert response.status_code == 404
    assert "尚未绑定" in response.json()["detail"]


def test_wechat_login_returns_token_for_bound_user(client, monkeypatch):
    from app.routers import auth as auth_router

    token = register(client, 21)

    async def fake_code2session(code: str) -> str:
        assert code in {"bind-code", "login-code"}
        return "bound-openid"

    monkeypatch.setattr(auth_router, "_code2session", fake_code2session)
    bound = client.post(
        "/api/auth/wechat/bind",
        headers=auth_header(token),
        json={"code": "bind-code"},
    )
    assert bound.status_code == 200, bound.text

    logged_in = client.post("/api/auth/wechat/login", json={"code": "login-code"})
    assert logged_in.status_code == 200, logged_in.text
    assert logged_in.json()["access_token"]
    assert logged_in.json()["user"]["student_id"] == "20260021"


def test_space_messages_require_usage_and_only_expose_display_name(client):
    outsider = register(client, 31)
    room_id = client.get("/api/rooms").json()[0]["id"]

    unauthorized = client.get(f"/api/rooms/{room_id}/messages")
    assert unauthorized.status_code in {401, 403}
    denied = client.post(
        f"/api/rooms/{room_id}/messages",
        headers=auth_header(outsider),
        json={"content": "尚未使用过这个空间", "photo_urls": []},
    )
    assert denied.status_code == 403
    assert "实际使用过" in denied.text

    user_token = register(client, 32)
    created = client.post("/api/reservations", headers=auth_header(user_token), json={
        "scene": "study",
        "date": (date.today() + timedelta(days=6)).isoformat(),
        "start_slot": 24,
        "end_slot": 25,
        "people_count": 1,
        "purpose": "个人自习",
        "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    })
    assert created.status_code == 201, created.text
    used_room_id = created.json()["room_id"]
    asyncio.run(force_used(created.json()["id"]))

    photo = client.post(
        f"/api/rooms/{used_room_id}/messages/photo",
        headers=auth_header(user_token),
        files={"file": ("space.jpg", b"\xff\xd8\xff\xe0message-photo", "image/jpeg")},
    )
    assert photo.status_code == 201, photo.text
    assert photo.json()["url"].startswith("/uploads/message_")

    posted = client.post(
        f"/api/rooms/{used_room_id}/messages",
        headers=auth_header(user_token),
        json={"content": "空间整洁，使用体验很好", "photo_urls": [photo.json()["url"]]},
    )
    assert posted.status_code == 201, posted.text
    assert posted.json()["author_name"] == "测试学生32"

    messages = client.get(
        f"/api/rooms/{used_room_id}/messages",
        headers=auth_header(user_token),
    )
    assert messages.status_code == 200, messages.text
    message = next(item for item in messages.json() if item["id"] == posted.json()["id"])
    assert message["content"] == "空间整洁，使用体验很好"
    assert message["photo_urls"] == [photo.json()["url"]]
    assert set(message) == {"id", "room_id", "content", "photo_urls", "author_name", "created_at"}

    empty = client.post(
        f"/api/rooms/{used_room_id}/messages",
        headers=auth_header(user_token),
        json={"content": "   ", "photo_urls": []},
    )
    assert empty.status_code == 422


def test_allocation_review_limit_priority_and_export(client):
    day = (date.today() + timedelta(days=1)).isoformat()
    student1 = register(client, 1)
    student2 = register(client, 2)
    student3 = register(client, 3)
    student4 = register(client, 4)

    booking = {"scene": "study", "date": day, "start_slot": 2, "end_slot": 4, "people_count": 1, "purpose": "个人自习", "campus_card_photo_url": "/uploads/campus_card_test.jpg"}
    first = client.post("/api/reservations", json=booking, headers=auth_header(student1))
    assert first.status_code == 201, first.text
    assert first.json()["room"]["room_code"] == "A102"
    assert first.json()["start_minute"] == 9 * 60
    assert first.json()["status"] == "pending"

    # Shared study capacity permits another one-person reservation in A102.
    second = client.post("/api/reservations", json=booking, headers=auth_header(student2))
    assert second.status_code == 201, second.text
    assert second.json()["room"]["room_code"] == "A102"

    # Existing 1h + another 3.5h exceeds the configurable 4-hour daily sum.
    over_limit = client.post("/api/reservations", json={**booking, "start_slot": 8, "end_slot": 15}, headers=auth_header(student1))
    assert over_limit.status_code == 400

    oversized_meeting = client.post("/api/reservations", headers=auth_header(student4), json={
        **booking,
        "scene": "meeting",
        "people_count": 500,
        "purpose": "召开书院学生工作协调会议",
    })
    assert oversized_meeting.status_code == 409

    admin_login = client.post("/api/auth/login", json={"student_id": "admin001", "password": "admin123"})
    assert admin_login.status_code == 200, admin_login.text
    admin = auth_header(admin_login.json()["access_token"])
    invalid_settings = client.put("/api/admin/settings", headers=admin, json={"values": {"open_hour": 8.5}})
    assert invalid_settings.status_code == 400
    assert "必须为整数" in invalid_settings.text
    reviewed = client.post("/api/admin/reservations/review", headers=admin, json={
        "reservation_ids": [first.json()["id"], second.json()["id"]], "decision": "approved", "note": "测试通过",
    })
    assert reviewed.json()["processed"] == 2

    music = client.post("/api/reservations", headers=auth_header(student2), json={
        "scene": "music", "date": day, "start_slot": 18, "end_slot": 20, "people_count": 1,
        "purpose": "进行个人钢琴练习训练", "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    })
    assert music.status_code == 201, music.text
    assert music.json()["room"]["room_code"] == "B102"
    piano = client.post("/api/reservations", headers=auth_header(student3), json={
        "scene": "music", "date": day, "start_slot": 18, "end_slot": 20, "people_count": 1,
        "purpose": "进行声乐与钢琴联合训练", "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    })
    assert piano.status_code == 201, piano.text
    assert piano.json()["room"]["room_code"] == "A103"
    event = client.post("/api/reservations", headers=auth_header(student4), json={
        "scene": "event", "date": day, "start_slot": 18, "end_slot": 20, "people_count": 40,
        "purpose": "举办书院大型主题宣讲活动", "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    })
    assert event.status_code == 409

    export = client.get("/api/admin/export.xlsx", headers=admin)
    assert export.status_code == 200
    assert export.content[:2] == b"PK"
    workbook = load_workbook(io.BytesIO(export.content), read_only=True)
    assert workbook.sheetnames == ["预约记录", "预约限制记录"]
    headers = [cell.value for cell in next(workbook["预约记录"].iter_rows())]
    assert "玉兰卡照片" in headers
    assert "违规标记" in headers

    # Closed-loop cleanup: upload -> admin reject + timed restriction -> booking blocked.
    asyncio.run(force_finished(first.json()["id"]))
    photo = client.post(
        "/api/reservations/photos", headers=auth_header(student1),
        files={"file": ("cleanup.jpg", b"\xff\xd8\xff\xe0test-jpeg-data", "image/jpeg")},
    )
    assert photo.status_code == 201, photo.text
    submitted = client.post(
        f"/api/reservations/{first.json()['id']}/cleanup",
        headers=auth_header(student1), json={"photo_urls": [photo.json()["url"]]},
    )
    assert submitted.status_code == 200, submitted.text
    # Requirement: submitting the cleanup photo immediately unlocks booking;
    # the administrator reviews it asynchronously afterwards.
    unlocked = client.post("/api/reservations", json={**booking, "date": (date.today() + timedelta(days=2)).isoformat()}, headers=auth_header(student1))
    assert unlocked.status_code == 201, unlocked.text
    queue = client.get("/api/admin/cleanup", headers=admin).json()
    cleanup_id = next(item["cleanup"]["id"] for item in queue if item["id"] == first.json()["id"])
    rejected = client.post(f"/api/admin/cleanup/{cleanup_id}/review", headers=admin, json={
        "decision": "rejected", "note": "照片无法确认清扫完成", "restrict_user": True,
        "restriction_level": "timed", "restriction_days": 7,
    })
    assert rejected.status_code == 200, rejected.text
    blocked = client.post("/api/reservations", json={**booking, "date": (date.today() + timedelta(days=2)).isoformat()}, headers=auth_header(student1))
    assert blocked.status_code == 403
    profile = client.get("/api/auth/me", headers=auth_header(student1)).json()
    assert profile["booking_restricted"] is True
    assert profile["restriction_level"] == "timed"
    assert "照片无法确认" in profile["restriction_reason"]


def test_approved_booking_auto_advances_without_checkin_or_no_show_penalty(client):
    token = register(client, 10)
    day = (date.today() + timedelta(days=1)).isoformat()
    created = client.post("/api/reservations", headers=auth_header(token), json={
        "scene": "study", "date": day, "start_slot": 4, "end_slot": 5,
        "people_count": 1, "purpose": "个人自习",
        "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    })
    assert created.status_code == 201, created.text

    admin_login = client.post("/api/auth/login", json={"student_id": "admin001", "password": "admin123"})
    admin = auth_header(admin_login.json()["access_token"])
    reviewed = client.post("/api/admin/reservations/review", headers=admin, json={
        "reservation_ids": [created.json()["id"]], "decision": "approved", "note": "测试通过",
    })
    assert reviewed.status_code == 200

    # Move the booking into the past. A single state refresh must go directly
    # to cleanup, without creating a missed/no-show violation.
    asyncio.run(force_approved_past(created.json()["id"]))

    mine = client.get("/api/reservations/my", headers=auth_header(token))
    assert mine.status_code == 200
    item = next(value for value in mine.json() if value["id"] == created.json()["id"])
    assert item["status"] == "cleanup_pending"
    users = client.get("/api/admin/users?search=20260010", headers=admin).json()
    violations = client.get(f"/api/admin/users/{users[0]['id']}/violations", headers=admin).json()
    assert violations == []


def test_three_cleanup_violations_trigger_thirty_day_ban(client):
    token = register(client, 11)
    admin_login = client.post("/api/auth/login", json={"student_id": "admin001", "password": "admin123"})
    admin = auth_header(admin_login.json()["access_token"])
    reservation_ids = []
    for index in range(3):
        payload = {
            "scene": "study",
            "date": (date.today() + timedelta(days=index + 1)).isoformat(),
            "start_slot": 4,
            "end_slot": 5,
            "people_count": 1,
            "purpose": "个人自习",
            "campus_card_photo_url": "/uploads/campus_card_test.jpg",
        }
        created = client.post("/api/reservations", json=payload, headers=auth_header(token))
        assert created.status_code == 201, created.text
        reservation_ids.append(created.json()["id"])
    reviewed = client.post("/api/admin/reservations/review", headers=admin, json={
        "reservation_ids": reservation_ids, "decision": "approved", "note": "测试通过",
    })
    assert reviewed.status_code == 200
    asyncio.run(force_cleanup_review(reservation_ids))
    cleanup_queue = client.get("/api/admin/cleanup", headers=admin).json()
    cleanup_ids = {
        item["id"]: item["cleanup"]["id"]
        for item in cleanup_queue
        if item["id"] in reservation_ids
    }
    assert set(cleanup_ids) == set(reservation_ids)
    for reservation_id in reservation_ids:
        rejected = client.post(
            f"/api/admin/cleanup/{cleanup_ids[reservation_id]}/review",
            headers=admin,
            json={"decision": "rejected", "note": "现场清扫核验不合格", "restrict_user": False},
        )
        assert rejected.status_code == 200, rejected.text

    # Use the admin API instead of relying on database-assigned user ids.
    users = client.get("/api/admin/users?search=20260011", headers=admin).json()
    user_id = users[0]["id"]
    violations = client.get(f"/api/admin/users/{user_id}/violations", headers=admin).json()
    restrictions = client.get(f"/api/admin/users/{user_id}/restrictions", headers=admin).json()
    assert len(violations) == 3
    assert {item["type"] for item in violations} == {"cleanup_failed"}
    automatic = next(item for item in restrictions if "自动禁约30天" in item["reason"])
    assert automatic["level"] == "timed"
    profile = client.get("/api/auth/me", headers=auth_header(token)).json()
    assert profile["booking_restricted"] is True
    assert profile["restriction_ends_at"] is not None
    assert "自动禁约30天" in profile["restriction_reason"]

    blocked = client.post("/api/reservations", json={
        "scene": "study", "date": (date.today() + timedelta(days=4)).isoformat(), "start_slot": 6, "end_slot": 7,
        "people_count": 1, "purpose": "个人自习", "campus_card_photo_url": "/uploads/campus_card_test.jpg",
    }, headers=auth_header(token))
    assert blocked.status_code == 403
    assert "自动禁约30天" in blocked.text
